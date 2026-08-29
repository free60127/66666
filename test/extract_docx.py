import sys
from docx import Document
names = ["lesson18","lesson20","lesson22"]
out = open(r"E:/AI/back translate/test/docx_dump.txt", "w", encoding="utf-8")
for name in names:
    path = r"E:/A新概念/新概念2/" + name + ".docx"
    try:
        doc = Document(path)
        out.write("="*20 + " " + name + " " + "="*20 + "\n")
        for i, p in enumerate(doc.paragraphs):
            t = p.text.strip()
            if t:
                out.write(f"[{i}][{p.style.name}] {t}\n")
        for ti, tbl in enumerate(doc.tables):
            out.write(f"--- table {ti} ---\n")
            for row in tbl.rows:
                cells = [c.text.strip().replace(chr(10),' / ') for c in row.cells]
                out.write(" | ".join(cells) + "\n")
    except Exception as e:
        out.write("ERR " + name + " " + repr(e) + "\n")
out.close()
print("done")
